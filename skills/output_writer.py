from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import uuid
from datetime import datetime
from pathlib import Path

from loguru import logger

from config import get_settings, get_user_profile
from schemas.models import ApplicationPackage, CoverLetterRenderDiagnostics


class OutputWriteError(RuntimeError):
    """Raised when an output package cannot be published atomically."""

    def __init__(self, stage: str, message: str, diagnostics_path: Path | None = None) -> None:
        self.stage = stage
        self.diagnostics_path = diagnostics_path
        super().__init__(f"output writing failed during {stage}: {message}")


def _slugify(text: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", text.strip().lower()).strip("_")
    return slug or "unknown"


def _output_name(company: str, role: str, session_id: str | None = None) -> str:
    now = datetime.now()
    date_part = now.strftime("%Y%m%d")
    c = _slugify(company)
    r = _slugify(role)
    run_suffix = _slugify(session_id or now.strftime("%H%M%S_%f"))[-24:]
    suffix = f"{c}_{r}"[:64]
    return f"{date_part}_{suffix}_{run_suffix}"


def _set_run_font(run, *, size: float, bold: bool = False, color: str = "111111") -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _add_line(document: Document, text: str, *, size: float = 10.5, bold: bool = False,
              color: str = "111111", after: float = 0) -> None:
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    _set_run_font(paragraph.add_run(text), size=size, bold=bold, color=color)


def _write_cover_letter_docx(package: ApplicationPackage, out_dir: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    profile = get_user_profile()
    settings = get_settings()
    candidate_name = str(profile.get("name") or "").strip()
    if not candidate_name:
        raise ValueError("candidate name is required for the cover-letter header and signature")
    today = datetime.now().strftime("%B %d, %Y")
    body = (package.cover_letter.body if package.cover_letter else "").strip()
    paragraphs = [p.strip() for p in body.split("\n\n") if p.strip()]
    if len(paragraphs) != 3:
        raise ValueError(f"expected exactly 3 body paragraphs, found {len(paragraphs)}")

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(0.9)
    section.header_distance = section.footer_distance = Inches(0.45)
    _add_line(document, candidate_name, size=18, bold=True, color="1F4D78", after=2)
    contact_groups = [
        [settings.cover_letter_email, settings.cover_letter_phone, settings.cover_letter_location],
        [settings.cover_letter_linkedin_url, settings.cover_letter_portfolio_url],
    ]
    known_groups = [[text for text in group if text] for group in contact_groups]
    known_groups = [group for group in known_groups if group]
    if known_groups:
        for group_index, known_contacts in enumerate(known_groups):
            contact_paragraph = document.add_paragraph()
            contact_paragraph.paragraph_format.space_before = Pt(0)
            contact_paragraph.paragraph_format.space_after = Pt(
                14 if group_index == len(known_groups) - 1 else 2
            )
            for index, text in enumerate(known_contacts):
                if index:
                    _set_run_font(contact_paragraph.add_run(" | "), size=9, color="777777")
                _set_run_font(contact_paragraph.add_run(text), size=9, color="555555")
    else:
        document.paragraphs[-1].paragraph_format.space_after = Pt(14)
    _add_line(document, today, after=10)

    recipient = package.cover_letter_recipient
    if recipient is not None:
        recipient_lines = [recipient.name, recipient.title, recipient.company, *recipient.address_lines]
        known_lines = [line for line in recipient_lines if line]
        for index, line in enumerate(known_lines):
            _add_line(document, line, after=8 if index == len(known_lines) - 1 else 0)

    for text in paragraphs:
        paragraph = document.add_paragraph()
        _set_run_font(paragraph.add_run(text), size=10.5)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(9)
        paragraph.paragraph_format.line_spacing = Pt(14)
    _add_line(document, "Sincerely,", after=12)
    _add_line(document, candidate_name, bold=True)

    path = out_dir / "cover_letter.docx"
    document.save(path)
    return path


def find_libreoffice() -> str | None:
    """Return the configured LibreOffice executable, if it is installed."""
    configured = get_settings().libreoffice_path or ""
    configured_cli = (
        str(Path(configured).with_suffix(".com"))
        if configured and Path(configured).suffix.casefold() == ".exe" and Path(configured).with_suffix(".com").is_file()
        else configured
    )
    candidates = [
        configured_cli,
        shutil.which("soffice.com") or "",
        shutil.which("soffice") or "",
        shutil.which("libreoffice") or "",
        r"C:\Program Files\LibreOffice\program\soffice.com",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.com",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    return next((candidate for candidate in candidates if candidate and Path(candidate).is_file()), None)


def find_pdftoppm() -> str | None:
    candidate = shutil.which("pdftoppm")
    return candidate if candidate and Path(candidate).is_file() else None


def _convert_docx_to_pdf(docx_path: Path, out_dir: Path) -> Path:
    soffice = find_libreoffice()
    if not soffice:
        raise FileNotFoundError("LibreOffice/soffice was not found; set LIBREOFFICE_PATH")
    out_dir = out_dir.resolve()
    docx_path = docx_path.resolve()
    with tempfile.TemporaryDirectory(prefix="caerus-lo-") as profile_name:
        profile_uri = Path(profile_name).resolve().as_uri()
        result = subprocess.run(
            [
                soffice,
                "--headless",
                "--norestore",
                "--nodefault",
                "--nofirststartwizard",
                "--nolockcheck",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(out_dir),
                str(docx_path),
            ],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
    pdf_path = out_dir / "cover_letter.pdf"
    if result.returncode != 0 or not pdf_path.is_file() or pdf_path.stat().st_size == 0:
        output = (result.stderr or result.stdout or "conversion produced no PDF").strip()
        detail = f"LibreOffice exited {result.returncode}: {output}"
        raise RuntimeError(detail[:500])
    return pdf_path


def _validate_cover_letter_artifacts(docx_path: Path, pdf_path: Path, expected_body: str) -> int:
    from docx import Document
    from pypdf import PdfReader

    reopened = Document(docx_path)
    docx_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    for paragraph in expected_body.split("\n\n"):
        if paragraph not in docx_text:
            raise ValueError("DOCX content does not contain the complete cover-letter body")
    reader = PdfReader(str(pdf_path))
    page_count = len(reader.pages)
    if page_count != 1:
        raise ValueError(f"cover-letter PDF must be exactly one page, found {page_count}")
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    expected_text = " ".join(re.findall(r"\w+", expected_body.casefold(), re.UNICODE))
    rendered_text = " ".join(re.findall(r"\w+", pdf_text.casefold(), re.UNICODE))
    if expected_text not in rendered_text:
        raise ValueError("PDF text validation failed")
    pdftoppm = find_pdftoppm()
    if not pdftoppm:
        raise FileNotFoundError("pdftoppm was not found; install Poppler for visual validation")
    preview_prefix = pdf_path.parent / ".cover_letter_preview"
    result = subprocess.run(
        [pdftoppm, "-singlefile", "-png", "-r", "100", str(pdf_path), str(preview_prefix)],
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )
    preview_path = preview_prefix.with_suffix(".png")
    if result.returncode != 0 or not preview_path.is_file():
        raise RuntimeError((result.stderr or "PDF rasterization failed").strip()[:500])
    from PIL import Image, ImageChops

    with Image.open(preview_path) as image:
        grayscale = image.convert("L")
        ink = ImageChops.invert(grayscale)
        bounds = ink.point(lambda value: 255 if value > 18 else 0).getbbox()
        if bounds is None:
            raise ValueError("rendered PDF page is blank")
        left, top, right, bottom = bounds
        safe_edge = 8
        if left <= safe_edge or top <= safe_edge or right >= image.width - safe_edge or bottom >= image.height - safe_edge:
            raise ValueError("rendered PDF content reaches the page edge and may be clipped")
    preview_path.unlink()
    return page_count


def _write_company_brief(package: ApplicationPackage, out_dir: Path) -> Path:
    path = out_dir / "company_brief.md"
    brief = package.company_brief
    lines = [
        f"# Company Brief: {brief.company}",
        f"- Research Status: {brief.research_status.value}",
        f"- Stage: {brief.stage.value}",
        f"- Fit Score: {brief.fit_score}",
        f"- Sponsorship: {brief.sponsorship}",
        "",
        "## Role Context",
        *([f"- {x}" for x in brief.role_context] or ["- Unknown"]),
        "",
        "## Recent Developments",
        *([f"- {x}" for x in brief.recent_developments] or ["- Unknown"]),
        "",
        "## Strong Overlaps",
        *([f"- {x}" for x in brief.strong_overlaps] or ["- Unknown"]),
        "",
        "## Candidate Overlaps",
        *([f"- {x}" for x in brief.candidate_overlaps] or ["- Unknown"]),
        "",
        "## Potential Angles",
        *([f"- {x}" for x in brief.potential_angles] or ["- Unknown"]),
        "",
        "## Tech Highlights",
        *([f"- {x}" for x in brief.tech_highlights] or ["- Unknown"]),
        "",
        "## Culture Notes",
        *([f"- {x}" for x in brief.culture_notes] or ["- Unknown"]),
        "",
        "## Talking Points",
        *([f"- {x}" for x in brief.talking_points] or ["- Unknown"]),
        "",
        "## Concerns or Unknowns",
        *([f"- {x}" for x in brief.concerns_or_unknowns] or ["- Unknown"]),
        "",
        "## Evidence",
        *(
            [f"- {claim.statement} [{', '.join(claim.source_ids)}]" for claim in brief.evidence]
            or ["- Unknown"]
        ),
        "",
        "## Sources",
        *(
            [
                f"- {source.id} — {source.title}"
                f"{f' ({source.published_date})' if source.published_date else ''} — {source.url}"
                for source in brief.sources
            ]
            or ["- Unknown"]
        ),
    ]
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return path


def _write_resume_report(package: ApplicationPackage, out_dir: Path) -> Path:
    path = out_dir / "resume_report.md"
    sel = package.resume_selection
    lines = [
        "# Resume Selection Report",
        f"- Variant: {sel.variant}",
        f"- Grade: {sel.grade}",
        f"- Fit Score: {sel.fit_score}",
        f"- Selected Resume Path: {sel.selected_resume_path}",
        "",
        "## Strengths",
        *[f"- {x}" for x in sel.strengths],
        "",
        "## Gaps",
        *[f"- {x}" for x in sel.gaps],
        "",
        "## Talking Points",
        *[f"- {x}" for x in sel.talking_points],
        "",
        "## Recommended Projects",
    ]
    for project in sel.project_recommendations:
        lines.extend(
            [
                f"### {project.name} ({project.score}/100)",
                f"- Project ID: {project.project_id}",
                f"- Repository: {project.repository_url or 'Not provided'}",
                f"- Source Ref: {project.source_ref or 'Not provided'}",
                f"- Why: {project.reason}",
                "- Score Breakdown: "
                f"technology {project.technology_score}/40, "
                f"domains {project.domain_score}/25, "
                f"company/role {project.company_role_score}/20, "
                f"tier {project.tier_score}/10, "
                f"recency {project.recency_score}/5",
                f"- Required Matches: {', '.join(project.matched_required) or 'None'}",
                f"- Preferred Matches: {', '.join(project.matched_preferred) or 'None'}",
                f"- Domain Matches: {', '.join(project.matched_domains) or 'None'}",
                f"- Company/Role Matches: {', '.join(project.matched_company_signals) or 'None'}",
                "",
            ]
        )
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return path


def write_outputs(package: ApplicationPackage) -> ApplicationPackage:
    settings = get_settings()
    company = package.jd.company or "unknown_company"
    role = package.jd.role or "unknown_role"
    outputs_root = Path(settings.outputs_dir)
    outputs_root.mkdir(parents=True, exist_ok=True)
    output_name = _output_name(company, role, package.session_id)
    out_dir = outputs_root / output_name
    if out_dir.exists():
        out_dir = outputs_root / f"{output_name}_{uuid.uuid4().hex[:8]}"
    diagnostics_root = outputs_root / ".diagnostics"
    diagnostics_root.mkdir(parents=True, exist_ok=True)
    staging_dir = diagnostics_root / f"{out_dir.name}.staging-{uuid.uuid4().hex[:8]}"
    staging_dir.mkdir()

    stage = "docx_generation"
    try:
        docx_path = _write_cover_letter_docx(package, staging_dir)
        stage = "pdf_conversion"
        pdf_path = _convert_docx_to_pdf(docx_path, staging_dir)
        stage = "artifact_validation"
        body = package.cover_letter.body if package.cover_letter else ""
        page_count = _validate_cover_letter_artifacts(docx_path, pdf_path, body)
        stage = "supporting_outputs"
        company_brief_path = _write_company_brief(package, staging_dir)
        resume_report_path = _write_resume_report(package, staging_dir)

        selected_name: str | None = None
        src = package.resume_selection.selected_resume_path if package.resume_selection else None
        if src:
            src_path = Path(src)
            if src_path.exists():
                selected_name = src_path.name
                shutil.copy2(src_path, staging_dir / selected_name)
            else:
                logger.warning("resume file missing, cannot copy: {}", src_path)

        stage = "publish"
        staging_dir.replace(out_dir)
    except Exception as exc:
        package.render_diagnostics = CoverLetterRenderDiagnostics(
            stage=stage,
            success=False,
            validation_codes=[stage],
            diagnostics_path=str(staging_dir),
            message=str(exc)[:500],
        )
        raise OutputWriteError(stage, str(exc), staging_dir) from exc

    package.output_dir = str(out_dir)
    package.cover_letter_path = str(out_dir / pdf_path.name)
    package.cover_letter_docx_path = str(out_dir / docx_path.name)
    package.company_brief_path = str(out_dir / company_brief_path.name)
    package.resume_report_path = str(out_dir / resume_report_path.name)
    package.selected_resume_copy_path = str(out_dir / selected_name) if selected_name else None
    package.render_diagnostics = CoverLetterRenderDiagnostics(
        stage="complete",
        success=True,
        pdf_page_count=page_count,
    )
    return package
